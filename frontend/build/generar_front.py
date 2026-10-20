import json, re, datetime
import openpyxl
SRC='/mnt/user-data/outputs/checklist_ley21719_integrado.xlsx'
wbf=openpyxl.load_workbook(SRC)            # fórmulas
def iso(v):
    if isinstance(v,datetime.datetime): return v.strftime('%Y-%m-%d')
    if isinstance(v,str):
        m=re.match(r'^(\d{2})/(\d{2})/(\d{4})$',v.strip())
        if m: return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return '' if v is None else str(v)
# controles
wc=wbf['Controles marcos']; ctrls=[]; rowkey={}
for r in range(2,wc.max_row+1):
    k=wc.cell(r,2).value
    if not k: continue
    rowkey[r]=k
    ctrls.append({"marco":wc.cell(r,1).value,"id":k,"descripcion":wc.cell(r,3).value,"tema":wc.cell(r,4).value})
# requisitos
wq=wbf['Requisitos Ley']; reqs=[]
for r in range(2,wq.max_row+1):
    iid=wq.cell(r,2).value
    if not iid: continue
    f=str(wq.cell(r,5).value or '')
    refs=[]
    for m in re.finditer(r"'Controles marcos'!\$E\$(\d+)",f):
        refs.append(rowkey[int(m.group(1))])
    for m in re.finditer(r"Resumen!\$C\$(\d)",f):
        refs.append({'4':'RAT','5':'ARCOP','6':'TRANSF'}[m.group(1)])
    reqs.append({"dim":wq.cell(r,1).value,"id":iid,"texto":wq.cell(r,3).value,
                 "controles":refs,"tipo":wq.cell(r,10).value,"incluir":wq.cell(r,11).value=="Sí"})
# dimensiones
wr=wbf['Resumen']; dims=[]
for r in range(1,wr.max_row+1):
    a=wr.cell(r,1).value
    if isinstance(a,str) and re.match(r'^D\d+$',a):
        dims.append({"codigo":a,"nombre":wr.cell(r,2).value,"peso":float(wr.cell(r,3).value)})
# tablas
def table(name, idprefix, nfields=None):
    ws=wbf[name]; headers=[str(c.value or '').replace('\n',' ').strip() for c in ws[1]]
    headers=[h for h in headers]
    rows=[]
    for r in range(2,ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1,len(headers)+1)]
        if isinstance(vals[0],str) and vals[0].startswith(idprefix):
            rows.append([iso(v) for v in vals])
    return headers, rows
rh,rr=table('RAT','RAT-')
ah,ar=table('ARCOP','ARC-')
th,tr=table('Transferencias','RAT-')
rh=rh[:16]; rr=[r[:16] for r in rr]
ah=ah[:8]; ar=[r[:8] for r in ar]
th=th[:9]; tr=[r[:9] for r in tr]
def schema(h):
    out=[]
    for i,l in enumerate(h):
        t='text'
        if i>0 and ('Fecha' in l or 'Última Revisión' in l): t='date'
        if l.startswith('Datos Sensibles'): t='select'
        out.append({"k":f"f{i}","label":l,"type":t})
    return out
sch={"rat":schema(rh),"arcop":schema(ah),"transf":schema(th)}

# --- campos E2E del RAT (se agregan al final del esquema) ---
NUEVOS_RAT=[
 ("Fuente de los datos","select",["Directamente del titular","Formulario o canal digital","Sistema interno","Tercero o proveedor","Fuente pública"]),
 ("Cómo se recolectan","text",None),
 ("Ubicación del almacenamiento","select",["Chile","Fuera de Chile","Chile y fuera de Chile","Por definir"]),
 ("Forma de eliminación","text",None),
 ("Encargados (proveedores que tratan datos)","text",None),
 ("Decisiones automatizadas","select",["No","Sí"]),
 ("Nivel de riesgo","select",["Bajo","Medio","Alto"]),
 ("Requiere evaluación de impacto (EIPD)","select",["No","Sí","Por evaluar"]),
 ("Derechos ARCOP que aplican","text",None),
 ("Estado de la actividad","select",["Vigente","En revisión","Dada de baja"]),
]
for n,(lab,tp,op) in enumerate(NUEVOS_RAT, start=16):
    e={"k":f"f{n}","label":lab,"type":"date" if False else ("select" if tp=="select" else "text")}
    if op: e["opts"]=op
    sch["rat"].append(e)
sch["rat"][5]["opts"]=["Sí","No"]
sch["rat"][3]["opts"]=["Lit. a — Consentimiento expreso","Lit. b — Ejecución de contrato","Lit. c — Obligación legal","Lit. d — Interés vital","Lit. e — Interés legítimo","Lit. f — Interés público o autoridad pública"]
sch["arcop"][2]["opts"]=["Acceso","Rectificación","Cancelación / Supresión","Oposición","Portabilidad","Decisiones automatizadas"]
sch["arcop"][6]["opts"]=["Pendiente","En proceso","Respondido","Rechazado"]
sch["arcop"][6]["type"]="select"
def rowsobj(rows):
    return [{f"f{i}":v for i,v in enumerate(r)} for r in rows]
seed={"dims":dims,"ctrls":ctrls,"reqs":reqs,"schemas":sch,
      "tables":{"rat":rowsobj(rr),"arcop":rowsobj(ar),"transf":rowsobj(tr)},
      "registros":[],"meta":{"org":"","responsable":"","fecha":"","evaluador":"","acepta":False,"aceptante":"","fechaAcepta":""}}


# --- hojas de referencia del Excel ---
we=wbf['Entrevistas']; guia=[]
for r in range(5,we.max_row+1):
    a=we.cell(r,1).value
    if isinstance(a,str) and re.match(r'^D\d+$',a):
        guia.append({"codigo":a,"dimension":we.cell(r,2).value,"entrevistado":we.cell(r,3).value,"preguntas":we.cell(r,4).value,"evidencia":we.cell(r,5).value})
wh=wbf['Cómo llenar']; como=[]
for r in range(6,wh.max_row+1):
    a=wh.cell(r,1).value
    if a is not None and str(a).strip().isdigit():
        como.append({"paso":str(a),"hoja":wh.cell(r,2).value,"que":wh.cell(r,3).value,"valores":wh.cell(r,4).value,"ejemplo":wh.cell(r,5).value})
wsc=wbf['Escala']; escala=[]
for r in range(2,wsc.max_row+1):
    if isinstance(wsc.cell(r,1).value,(int,float)):
        escala.append({"puntaje":wsc.cell(r,1).value,"nivel":wsc.cell(r,2).value,"criterio":wsc.cell(r,3).value})
wn=wbf['Notas']; notas=[wn.cell(r,1).value for r in range(1,wn.max_row+1) if wn.cell(r,1).value]
exec(open('/tmp/catrat.py',encoding='utf-8').read())
seed.update({"catRat":catrat,"guia":guia,"comoLlenar":como,"escala":escala,"notas":notas})

html=open('/tmp/front_template.html',encoding='utf-8').read()
assert '__SEED__' in html
html=html.replace('__SEED__', json.dumps(seed,ensure_ascii=False))
open('/mnt/user-data/outputs/checklist_ley21719_frontend.html','w',encoding='utf-8').write(html)
print('dims',len(dims),'ctrls',len(ctrls),'reqs',len(reqs),'rat',len(rr),'arc',len(ar),'tr',len(tr))
