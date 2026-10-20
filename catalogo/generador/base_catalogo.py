from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NAVY="1F3864"; LIGHT="D9E2F3"; INPUT="FFF2CC"
thin=Side(style="thin",color="BFBFBF"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor=NAVY)
WR=Alignment(wrap_text=True,vertical="top"); CT=Alignment(horizontal="center",vertical="center",wrap_text=True)

DIMS=[
("D1","Gobierno y responsabilidad proactiva",9),
("D2","Inventario y Registro de Actividades de Tratamiento (RAT)",9),
("D3","Licitud del tratamiento y consentimiento",9),
("D4","Derechos de los titulares (ARCOP)",9),
("D5","Principios y ciclo de vida del dato",7),
("D6","Datos sensibles, biométricos y de menores",6),
("D7","Evaluación de impacto y privacidad desde el diseño",6),
("D8","Seguridad de la información",11),
("D9","Gestión de incidentes y notificación de vulneraciones",9),
("D10","Encargados, terceros y transferencias internacionales",7),
("D11","Capacitación y cultura",4),
("D12","Auditoría, monitoreo y mejora continua",4),
("D13","Marcos de referencia complementarios (NIST CSF 2.0, CIS, COBIT, PCI DSS, normativa)",10),
]

ITEMS={
"D1":[
("1.1","Política de protección de datos personales aprobada por la dirección, publicada y comunicada al personal."),
("1.2","Responsable de cumplimiento (DPO o equivalente) designado formalmente, con funciones y autoridad definidas."),
("1.3","Recursos humanos, técnicos y presupuestarios asignados al programa de cumplimiento."),
("1.4","Programa de adecuación o modelo de prevención de infracciones documentado."),
("1.5","Evidencias de cumplimiento organizadas y disponibles para ser presentadas en cualquier momento (accountability)."),
("1.6","Rol de la organización (responsable o encargado) determinado para cada tratamiento."),
],
"D2":[
("2.1","RAT con finalidad, base de licitud, categorías de datos, titulares, destinatarios, plazos de conservación y medidas de seguridad."),
("2.2","Inventario de sistemas, aplicaciones y bases de datos que tratan datos personales."),
("2.3","Mapa de flujos de datos internos y externos, incluyendo nube y proveedores."),
("2.4","RAT revisado al menos anualmente y ante cambios significativos de procesos o sistemas."),
("2.5","Tratamientos de alto riesgo identificados y marcados en el RAT."),
],
"D3":[
("3.1","Cada tratamiento tiene una base de licitud identificada y documentada antes de iniciar."),
("3.2","Consentimiento libre, informado, específico e inequívoco, sin casillas premarcadas."),
("3.3","Registro de consentimientos con fecha, canal, versión del aviso y finalidad aceptada."),
("3.4","Revocación del consentimiento disponible con facilidad equivalente a su otorgamiento."),
("3.5","Ponderación de interés legítimo documentada cuando ese sea la base de licitud."),
("3.6","Información al titular entregada al momento de la recolección (aviso de privacidad completo)."),
],
"D4":[
("4.1","Procedimiento formal para atender solicitudes ARCOP, con canal publicado."),
("4.2","Plazos de respuesta definidos y monitoreados."),
("4.3","Verificación de identidad del solicitante antes de entregar o modificar datos."),
("4.4","Capacidad técnica para ubicar, rectificar, suprimir, oponerse y exportar datos en todos los sistemas."),
("4.5","Registro de solicitudes, respuestas y motivos de rechazo."),
("4.6","Tratamiento de decisiones automatizadas con derecho a información y revisión humana, cuando aplique."),
],
"D5":[
("5.1","Finalidades declaradas y limitación de uso posterior para otra finalidad incompatible."),
("5.2","Minimización: solo se recolectan y mantienen datos estrictamente necesarios."),
("5.3","Mecanismos para mantener los datos exactos y actualizados."),
("5.4","Política de retención con plazos concretos por categoría de dato."),
("5.5","Eliminación o anonimización segura al cumplirse el plazo, con evidencia."),
("5.6","Datos obtenidos de terceros o fuentes públicas validados en licitud y calidad antes de usarse."),
],
"D6":[
("6.1","Datos sensibles, biométricos, de salud y de menores identificados en el inventario."),
("6.2","Consentimiento explícito para el tratamiento de datos sensibles."),
("6.3","Controles reforzados para datos sensibles (acceso restringido, cifrado, logging)."),
("6.4","Tratamiento de datos de niños, niñas y adolescentes con interés superior y consentimiento de representante cuando corresponda."),
("6.5","Restricciones para tratamientos de datos biométricos y de vigilancia sistemática."),
],
"D7":[
("7.1","Criterios de alto riesgo definidos y comunicados a las áreas que proponen tratamientos."),
("7.2","EIPD/DPIA documentada antes de iniciar cada tratamiento de alto riesgo."),
("7.3","Privacidad desde el diseño incorporada como control en gestión de proyectos y cambios."),
("7.4","Configuraciones por defecto de sistemas y productos con el mínimo de datos y accesos."),
("7.5","Procedimiento de consulta previa a la autoridad cuando el riesgo residual sea alto."),
],
"D8":[
("8.1","Sistema de gestión de seguridad de la información (ISO/IEC 27001) implementado o en implementación."),
("8.2","Control de acceso con mínimo privilegio, separación de funciones y revisión periódica de accesos."),
("8.3","Autenticación multifactor (MFA) en todos los sistemas que contienen datos personales."),
("8.4","Cifrado de datos en tránsito y en reposo, incluyendo respaldos y equipos portátiles."),
("8.5","Gestión de vulnerabilidades con parcheo según criticidad y plazos definidos."),
("8.6","Registros de auditoría de accesos a datos personales, protegidos y con retención definida."),
("8.7","Respaldos cifrados con pruebas de restauración documentadas."),
("8.8","Pruebas de intrusión o evaluaciones técnicas al menos anuales, con hallazgos críticos cerrados."),
("8.9","Seudonimización o anonimización aplicada cuando la finalidad lo permite."),
("8.10","Cuentas privilegiadas gestionadas con control reforzado y registro de sesiones."),
],
"D9":[
("9.1","Procedimiento de respuesta a incidentes con roles y responsabilidades asignados."),
("9.2","Mecanismo de detección y registro de vulneraciones que afecten datos personales."),
("9.3","Notificación a la autoridad dentro del plazo legal (ver nota sobre plazo en la hoja Notas)."),
("9.4","Notificación a los titulares afectados cuando exista riesgo elevado para ellos."),
("9.5","Plantilla de notificación a la autoridad y a titulares lista para uso."),
("9.6","Simulacro de brecha con datos personales realizado al menos una vez al año."),
("9.7","Bitácora de vulneraciones con causa raíz, impacto y acciones correctivas."),
("9.8","Plan de continuidad y recuperación ante desastres con RTO y RPO definidos y probados."),
],
"D10":[
("10.1","Inventario de encargados y proveedores que tratan datos personales por cuenta de la organización."),
("10.2","Contratos de encargo con las cláusulas exigidas: objeto, duración, obligaciones, subencargo y eliminación al término."),
("10.3","Evaluación de seguridad y privacidad de proveedores antes de contratar y periódicamente."),
("10.4","Derecho de auditoría sobre encargados establecido contractualmente."),
("10.5","Transferencias internacionales solo a países con nivel adecuado o con garantías contractuales."),
("10.6","Acuerdos de confidencialidad con terceros que accedan a información."),
],
"D11":[
("11.1","Capacitación en protección de datos al ingreso y con frecuencia anual."),
("11.2","Contenidos diferenciados por rol (TI, ventas, RRHH, atención al cliente, dirección)."),
("11.3","Campañas de concienciación y simulacros de phishing con métricas."),
("11.4","Compromisos de confidencialidad firmados por todo el personal y contratistas."),
],
"D12":[
("12.1","Plan anual de auditoría interna de cumplimiento aprobado."),
("12.2","Indicadores de privacidad (solicitudes ARCOP, incidentes, capacitación, EIPD) medidos y reportados a la dirección."),
("12.3","Hallazgos de auditoría con plan de acción, responsable y plazo."),
("12.4","Revisión anual de la política y del programa de cumplimiento."),
("12.5","Monitoreo de cambios normativos, reglamentos y criterios de la autoridad de control."),
],
"D13":[
("13.1","Control NIST ID.RA-01 aplicado: vulnerabilidades de sistemas con datos personales identificadas, validadas y registradas."),
("13.2","Control NIST ID.RA-08 aplicado: canal y proceso para recibir y responder divulgaciones de vulnerabilidades."),
("13.3","Controles NIST PR.PS-01 y CIS 4 aplicados: configuración segura documentada en sistemas con datos personales."),
("13.4","Controles NIST PR.PS-02 y PR.PS-03 aplicados: software y hardware sin soporte retirados o aislados si contienen datos personales."),
("13.5","Controles NIST PR.PS-05 y CIS 2 aplicados: software no autorizado bloqueado y registrado."),
("13.6","Control CIS 9 aplicado: protección de correo y navegador frente a phishing y sitios maliciosos."),
("13.7","Control CIS 10 aplicado: defensas anti-malware en todos los endpoints y servidores."),
("13.8","Controles NIST RS.AN-06 y RS.AN-07 aplicados: acciones y datos de incidentes registrados con integridad y procedencia."),
("13.9","Preservación de evidencia electrónica conforme a Ley 21.459 (arts. 14 y 18 N°1) incorporada al procedimiento de incidentes."),
("13.10","Controles NIST RS.MI-01 y RS.MI-02 aplicados: contención y erradicación documentadas para incidentes con datos personales."),
("13.11","Controles NIST RC.RP-03 y CIS 11 aplicados: integridad de respaldos verificada antes de restaurar datos personales."),
("13.12","Controles COBIT BAI06 y BAI07 aplicados: cambios con evaluación de impacto en privacidad y aceptación formal."),
("13.13","Controles COBIT BAI09 y BAI10 aplicados: activos y configuraciones gestionados en todo su ciclo de vida."),
("13.14","Benchmarks BM-OS, BM-DB, BM-CLOUD, BM-NET, BM-APP y BM-M365 aplicados a sistemas con datos personales."),
("13.15","Requisitos PCI DSS 2 y 5 aplicados al entorno de pagos (N/A si no se procesan tarjetas)."),
("13.16","Requisitos PCI DSS 7 y 8 extendidos a sistemas de pago y a sus usuarios (N/A si no se procesan tarjetas)."),
("13.17","Modelo de prevención de delitos (Ley 20.393, art. 4) integrado con el programa de datos personales."),
("13.18","Alineación con la Política Nacional de Ciberseguridad (PNCS) documentada en el programa de cumplimiento."),
],
}

wb=Workbook()

# Portada
ws=wb.active; ws.title="Portada"
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=38; ws.column_dimensions["C"].width=70
ws["B2"]="CHECKLIST DE CUMPLIMIENTO CON SCORING — LEY N° 21.719"; ws["B2"].font=Font(bold=True,size=14,color=NAVY)
ws["B3"]="Protección y tratamiento de datos personales — Chile"; ws["B3"].font=Font(italic=True)
info=[("Organización","[Completar]"),("Responsable de cumplimiento","[Completar]"),("Fecha de evaluación","[Completar]"),("Próxima revisión","[Completar]"),("Evaluador","[Completar]")]
r=5
for k,v in info:
    ws.cell(r,2,k).font=Font(bold=True); c=ws.cell(r,3,v); c.fill=PatternFill("solid",fgColor=INPUT); r+=1
r+=1
ws.cell(r,2,"CÓMO USAR ESTE CHECKLIST").font=Font(bold=True,color=NAVY); r+=1
steps=["1. Completar la hoja 'Checklist' asignando a cada requisito un puntaje de 1 a 5 (o N/A si no aplica, con justificación).",
"2. Registrar la evidencia concreta que respalda cada puntaje en la columna 'Evidencia'.",
"3. La hoja 'Resumen' calcula automáticamente el puntaje por dimensión y el total ponderado.",
"4. Los pesos por dimensión están en la hoja 'Resumen' (celda editable) y suman 100%.",
"5. Los requisitos sin puntaje no se consideran en los promedios, por lo que el total solo refleja lo evaluado."]
for s in steps:
    ws.cell(r,2,s); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); ws.cell(r,2).alignment=WR; ws.row_dimensions[r].height=30; r+=1

# Escala
ws2=wb.create_sheet("Escala de puntaje")
ws2.column_dimensions["A"].width=10; ws2.column_dimensions["B"].width=22; ws2.column_dimensions["C"].width=80
ws2.append(["Puntaje","Nivel","Criterio"])
scale=[(1,"No existe","No hay control ni práctica documentada."),
(2,"Inicial","Existe una práctica informal o parcial, sin documentación ni responsable definido."),
(3,"Definido","Control documentado y con responsable, pero aplicado de forma inconsistente o sin evidencia completa."),
(4,"Implementado","Control implementado y aplicado en la operación, con evidencia disponible."),
(5,"Optimizado","Control implementado, evidenciado, medido y revisado periódicamente con mejora continua.")]
for row in scale: ws2.append(list(row))
ws2.append([]); ws2.append(["Nivel de madurez del total","Rango","Interpretación"])
bands=[("Crítico","1.0 – 1.9","Exposición alta a sanciones. Plan de acción inmediato."),
("Inicial","2.0 – 2.9","Brechas significativas. Priorizar controles de D1, D2, D3 y D9."),
("En desarrollo","3.0 – 3.9","Base instalada con brechas. Cerrar evidencias y consistencia."),
("Avanzado","4.0 – 4.5","Cumplimiento sólido. Enfocarse en medición y auditoría."),
("Optimizado","4.6 – 5.0","Cumplimiento maduro con mejora continua demostrable.")]
for b in bands: ws2.append(list(b))
for c in ws2[1]+ws2[7]: c.font=H; c.fill=HF
for row in ws2.iter_rows(min_row=2):
    for c in row: c.alignment=WR; c.border=B

# Checklist
wc=wb.create_sheet("Checklist")
headers=["Dim.","Dimensión","ID","Requisito","Referencia orientativa","Puntaje (1-5 / N/A)","Evidencia requerida / disponible","Brecha identificada","Plan de acción","Responsable","Fecha límite"]
wc.append(headers)
refs={"D1":"Principio de responsabilidad proactiva; programa de cumplimiento","D2":"RAT (Ley 21.719, art. 14 bis según texto de la política; verificar)",
"D3":"Art. 12 y 13 (consentimiento y bases de licitud; verificar)","D4":"Derechos ARCOP (verificar plazos en texto oficial)",
"D5":"Principios de finalidad, minimización, calidad y conservación","D6":"Régimen de datos sensibles y de niños (verificar)",
"D7":"EIPD para alto riesgo (art. 15 ter según política; verificar)","D8":"Deber de seguridad de los datos; ISO/IEC 27001:2022 Anexo A",
"D9":"Reporte de vulneraciones (verificar plazo y destinatario)","D10":"Encargo del tratamiento y transferencias internacionales (verificar)",
"D11":"Medidas organizativas de seguridad","D12":"Responsabilidad proactiva y supervisión",
"D13":"Controles del alcance Canontex (IDs del txt); evaluar aplicabilidad"}
row=2
for code,name,_ in DIMS:
    for iid,req in ITEMS[code]:
        wc.append([code,name,iid,req,refs[code],None,None,None,None,None,None])
        row+=1
last=row-1
for col,w in zip("ABCDEFGHIJK",[7,30,7,60,34,16,38,34,34,16,14]):
    wc.column_dimensions[col].width=w
for c in wc[1]: c.font=H; c.fill=HF; c.alignment=CT; c.border=B
for rr in wc.iter_rows(min_row=2,max_row=last):
    for c in rr:
        c.alignment=WR; c.border=B
    rr[5].fill=PatternFill("solid",fgColor=INPUT); rr[5].alignment=CT
wc.freeze_panes="E2"; wc.auto_filter.ref=f"A1:K{last}"
dv=DataValidation(type="list",formula1='"1,2,3,4,5,N/A"',allow_blank=True)
dv.error="Ingrese 1, 2, 3, 4, 5 o N/A"; wc.add_data_validation(dv); dv.add(f"F2:F{last}")
wc.conditional_formatting.add(f"F2:F{last}",CellIsRule(operator="between",formula=["1","2"],fill=PatternFill("solid",fgColor="F8CBAD")))
wc.conditional_formatting.add(f"F2:F{last}",CellIsRule(operator="equal",formula=["3"],fill=PatternFill("solid",fgColor="FFE699")))
wc.conditional_formatting.add(f"F2:F{last}",CellIsRule(operator="between",formula=["4","5"],fill=PatternFill("solid",fgColor="C6E0B4")))

# Resumen
wr=wb.create_sheet("Resumen")
wr["A1"]="RESUMEN DE SCORING — LEY 21.719"; wr["A1"].font=Font(bold=True,size=13,color=NAVY)
hdr=["Código","Dimensión","Peso","Requisitos","Evaluados","Puntaje (1-5)","% cumplimiento","Nivel"]
for i,h in enumerate(hdr,1):
    c=wr.cell(3,i,h); c.font=H; c.fill=HF; c.alignment=CT; c.border=B
widths=[10,56,10,12,12,14,16,16]
for i,w in enumerate(widths,1): wr.column_dimensions[chr(64+i)].width=w
R="Checklist!"
first=4
for idx,(code,name,w) in enumerate(DIMS):
    r=first+idx
    wr.cell(r,1,code); wr.cell(r,2,name)
    wr.cell(r,3,w/100).number_format="0%"
    wr.cell(r,3).fill=PatternFill("solid",fgColor=INPUT)
    wr.cell(r,4,f'=COUNTIF({R}$A$2:$A${last},A{r})')
    wr.cell(r,5,f'=COUNTIFS({R}$A$2:$A${last},A{r},{R}$F$2:$F${last},">0")')
    wr.cell(r,6,f'=IFERROR(AVERAGEIFS({R}$F$2:$F${last},{R}$A$2:$A${last},A{r},{R}$F$2:$F${last},">0"),0)').number_format="0.00"
    wr.cell(r,7,f'=IF(F{r}>0,F{r}/5,"")').number_format="0%"
    wr.cell(r,8,f'=IF(F{r}=0,"Sin evaluar",IF(F{r}<2,"Crítico",IF(F{r}<3,"Inicial",IF(F{r}<4,"En desarrollo",IF(F{r}<4.6,"Avanzado","Optimizado")))))')
    for i in range(1,9): wr.cell(r,i).border=B
lr=first+len(DIMS)-1
tr=lr+1
wr.cell(tr,1,"TOTAL").font=Font(bold=True)
wr.cell(tr,2,"Puntaje total ponderado (solo dimensiones evaluadas)").font=Font(bold=True)
wr.cell(tr,3,f"=SUM(C{first}:C{lr})").number_format="0%"
wr.cell(tr,4,f"=SUM(D{first}:D{lr})")
wr.cell(tr,5,f"=SUM(E{first}:E{lr})")
wr.cell(tr,6,f'=IFERROR(SUMPRODUCT((F{first}:F{lr}>0)*F{first}:F{lr}*C{first}:C{lr})/SUMPRODUCT((F{first}:F{lr}>0)*C{first}:C{lr}),0)').number_format="0.00"
wr.cell(tr,7,f'=IF(F{tr}>0,F{tr}/5,"")').number_format="0%"
wr.cell(tr,8,f'=IF(F{tr}=0,"Sin evaluar",IF(F{tr}<2,"Crítico",IF(F{tr}<3,"Inicial",IF(F{tr}<4,"En desarrollo",IF(F{tr}<4.6,"Avanzado","Optimizado")))))').font=Font(bold=True)
for i in range(1,9):
    wr.cell(tr,i).border=B; wr.cell(tr,i).fill=PatternFill("solid",fgColor=LIGHT)
wr.cell(tr+1,1,f'=IF(ROUND(C{tr},4)=1,"Pesos correctos (100%)","ATENCIÓN: los pesos no suman 100%")')
wr.cell(tr+1,1).font=Font(italic=True,color="595959")
wr.cell(tr+3,1,"Lectura: el puntaje se calcula como promedio de los requisitos evaluados de cada dimensión, y el total pondera las dimensiones según su peso. Los requisitos sin puntaje no cuentan.").alignment=WR
wr.merge_cells(start_row=tr+3,start_column=1,end_row=tr+3,end_column=8); wr.row_dimensions[tr+3].height=30
wr.freeze_panes="A4"

# Notas
wn=wb.create_sheet("Notas")
wn.column_dimensions["A"].width=120
notes=["NOTAS DE VERIFICACIÓN (leer antes de usar)",
"1. Vigencia: la Ley 21.719 entra en vigencia el 1 de diciembre de 2026. Existe un proyecto para postergarla al 1 de diciembre de 2027; verificar su estado en el Congreso antes de cada evaluación.",
"2. Plazo de notificación de vulneraciones: las fuentes consultadas no coinciden. Algunas indican 72 horas, otras señalan que la ley no fija horas y exige notificar sin dilación indebida. Confirmar el texto oficial en la BCN antes de fijar plazos internos. El requisito 9.3 está redactado para tomar el plazo que resulte del texto.",
"3. Referencias a artículos: las citas de la columna 'Referencia orientativa' provienen de los documentos del proyecto (política de seguridad, checklist existente, matriz de cumplimiento) y no han sido verificadas una por una contra el texto oficial. Validar antes de usarlas en un informe formal.",
"4. Sanciones: el checklist anterior del proyecto indica montos en UTM y un rango de artículos (36-38) que no pude confirmar; las fuentes consultadas señalan cifras distintas para el máximo. No se incluyen montos en este archivo.",
"5. Delegado de protección de datos: el alcance de la obligación según el tamaño y tipo de tratamiento debe confirmarse en el texto oficial. El requisito 1.2 se mantiene como buena práctica y requisito de gobierno.",
"6. D13 referencia los controles del archivo de alcance por ID, sin repetir sus descripciones. PCI DSS y COBIT aplican solo si la organización procesa tarjetas o gestiona esos procesos; usar N/A con justificación cuando corresponda.",
"7. Este archivo es una herramienta de autoevaluación. No constituye asesoría legal; para una conclusión jurídica definitiva, validar con abogado especialista."]
for i,n in enumerate(notes,1):
    c=wn.cell(i,1,n); c.alignment=WR
    if i==1: c.font=Font(bold=True,color=NAVY,size=12)
    else: wn.row_dimensions[i].height=45

wb.save("/mnt/user-data/outputs/checklist_ley21719_scoring.xlsx")
print("ok", last-1, "items")
